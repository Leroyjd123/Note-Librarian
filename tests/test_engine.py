"""Engine tests: model output is validated before it can touch the workbook."""
import asyncio
import json
import os
import sys
from types import SimpleNamespace

import openpyxl
import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from backend import engine  # noqa: E402
from backend.providers.base import LLMProvider, ModelDeclined  # noqa: E402
from backend.xlsx_io import WorkbookEditor  # noqa: E402


def _workbook(path, rows=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Notes"
    ws.append(["TITLE", "NOTE", "TAGS", "COLOR"])
    for r in rows or [["Old title", "Old note", "Faith | Prayer", 4]]:
        ws.append(r)
    wb.save(path)


def _apply(path, results, allowed=None, mode="full"):
    wb = WorkbookEditor(path, backup=False)
    out = engine._apply(wb, wb.header_map(), mode, results, None, allowed)
    wb.save()
    return out, openpyxl.load_workbook(path)["Notes"]


def test_missing_tags_key_keeps_existing_tags(tmp_path):
    path = str(tmp_path / "wb.xlsx")
    _workbook(path)
    _, ws = _apply(path, [{"row": 2, "color": 4, "confidence": "HIGH"}], {2})
    assert ws["C2"].value == "Faith | Prayer"


def test_list_tags_are_joined_not_crashing(tmp_path):
    path = str(tmp_path / "wb.xlsx")
    _workbook(path)
    out, ws = _apply(path, [{"row": 2, "tags": ["Faith", "Jehovah"]}], {2})
    assert ws["C2"].value == "Faith | Jehovah"
    assert out["tags"] == 1


def test_rows_outside_batch_are_ignored(tmp_path):
    path = str(tmp_path / "wb.xlsx")
    _workbook(path)
    out, ws = _apply(
        path, [{"row": 1, "title": "Oops", "title_changed": True, "tags": "X"}], {2}
    )
    assert [c.value for c in ws[1]] == ["TITLE", "NOTE", "TAGS", "COLOR"]
    assert out["ignored"] == ["1"]


def test_duplicate_row_results_apply_once(tmp_path):
    path = str(tmp_path / "wb.xlsx")
    _workbook(path)
    out, ws = _apply(
        path,
        [
            {"row": 2, "title": "First good title here", "title_changed": True},
            {"row": 2, "title": "Second title", "title_changed": True},
        ],
        {2},
    )
    assert ws["A2"].value == "First good title here"
    assert out["ignored"] == ["2"]


@pytest.mark.parametrize("bad", [0, 7, 3.7, "x", True, None])
def test_out_of_range_colour_is_dropped(tmp_path, bad):
    path = str(tmp_path / "wb.xlsx")
    _workbook(path)
    _, ws = _apply(path, [{"row": 2, "color": bad}], {2})
    assert ws["D2"].value == 4


@pytest.mark.parametrize("good,expected", [(2, 2), ("5", 5), (6.0, 6)])
def test_valid_colour_is_written(tmp_path, good, expected):
    path = str(tmp_path / "wb.xlsx")
    _workbook(path)
    _, ws = _apply(path, [{"row": 2, "color": good}], {2})
    assert ws["D2"].value == expected


def test_non_string_title_is_ignored(tmp_path):
    path = str(tmp_path / "wb.xlsx")
    _workbook(path)
    _, ws = _apply(path, [{"row": 2, "title": {"x": 1}, "title_changed": True}], {2})
    assert ws["A2"].value == "Old title"


# ----- run_job -----

class FakeProvider(LLMProvider):
    """Answers each batch with ``reply(rows)``; records every call."""

    def __init__(self, reply):
        super().__init__("key", "model", 1000)
        self.reply = reply
        self.calls = 0

    def complete(self, system, user):
        self.calls += 1
        batch = json.loads(user.split("\n\n", 1)[1])
        return self.reply([b["row"] for b in batch])


def _req(**kw):
    base = dict(
        provider="anthropic", model="m", api_key="k", mode="full", note_level="reconstruct",
        batch_size=2, concurrency=1, max_tokens=1000, sheet_name=None,
        write_type=False, write_confidence=False, write_review=False,
    )
    base.update(kw)
    return SimpleNamespace(**base)


def _job(path):
    return SimpleNamespace(
        status="pending", total_notes=0, processed_notes=0, batches_total=0, batches_done=0,
        cells_written=0, title_edits=0, note_edits=0, tag_edits=0, color_edits=0,
        review=[], pending={}, warnings=[], error=None, output_path=path,
    )


def _run(monkeypatch, path, provider, **kw):
    monkeypatch.setattr(engine, "get_provider", lambda *a: provider)
    monkeypatch.setattr(engine, "RETRY_BASE_DELAY", 0)
    job = _job(path)
    asyncio.run(engine.run_job(job, path, _req(**kw)))
    return job


ROWS = [[f"T{i}", f"N{i}", "Faith", 4] for i in range(2, 7)]  # sheet rows 2..6


def test_missing_results_are_warned(tmp_path, monkeypatch):
    path = str(tmp_path / "wb.xlsx")
    _workbook(path, ROWS)
    # Return results for all but the last row of each batch (>= 60%, so accepted).
    def reply(rows):
        keep = rows if len(rows) == 1 else rows[:-1]
        return json.dumps({"results": [{"row": r, "color": 1} for r in keep]})
    job = _run(monkeypatch, path, FakeProvider(reply), batch_size=5)
    assert job.status == "done"
    assert any("No result returned for rows 6" in w for w in job.warnings)


def test_failed_job_writes_nothing_and_stops(tmp_path, monkeypatch):
    path = str(tmp_path / "wb.xlsx")
    _workbook(path, ROWS)
    with open(path, "rb") as f:
        before = f.read()

    class Boom(Exception):
        status_code = 401

    def reply(rows):
        if rows[0] >= 4:
            raise Boom("bad key")
        return json.dumps({"results": [{"row": r, "color": 1} for r in rows]})

    provider = FakeProvider(reply)
    job = _run(monkeypatch, path, provider, batch_size=2, concurrency=1)
    assert job.status == "error" and "bad key" in job.error
    with open(path, "rb") as f:
        assert f.read() == before  # earlier successful batch was not saved
    assert provider.calls == 2  # the third batch never ran


def test_transient_errors_are_retried(tmp_path, monkeypatch):
    path = str(tmp_path / "wb.xlsx")
    _workbook(path)

    class RateLimitError(Exception):
        status_code = 429

    attempts = {"n": 0}

    def reply(rows):
        attempts["n"] += 1
        if attempts["n"] < 3:
            raise RateLimitError("slow down")
        return json.dumps({"results": [{"row": 2, "color": 2}]})

    job = _run(monkeypatch, path, FakeProvider(reply))
    assert job.status == "done", job.error
    assert attempts["n"] == 3
    assert openpyxl.load_workbook(path)["Notes"]["D2"].value == 2


def test_declined_batch_is_split_not_fatal(tmp_path, monkeypatch):
    path = str(tmp_path / "wb.xlsx")
    _workbook(path, ROWS[:2])

    def reply(rows):
        if 3 in rows:
            raise ModelDeclined("declined")
        return json.dumps({"results": [{"row": r, "color": 3} for r in rows]})

    job = _run(monkeypatch, path, FakeProvider(reply))
    assert job.status == "done"
    ws = openpyxl.load_workbook(path)["Notes"]
    assert ws["D2"].value == 3 and ws["D3"].value == 4
    assert any("Row 3" in w for w in job.warnings)
    assert not any("No result returned" in w for w in job.warnings)  # not warned twice


def test_blank_rows_are_not_sent_or_written(tmp_path, monkeypatch):
    path = str(tmp_path / "wb.xlsx")
    _workbook(path, [["t2", "n2", "Faith", 4], [None, None, None, None], ["t4", "", "", 4]])
    sent = []

    def reply(rows):
        sent.extend(rows)
        return json.dumps({"results": [{"row": r, "tags": "Stub", "color": 1} for r in rows]})

    job = _run(monkeypatch, path, FakeProvider(reply), write_type=True)
    assert job.status == "done"
    assert sent == [2, 4]  # row 3 is blank; row 4 has a title so it is still processed
    ws = openpyxl.load_workbook(path)["Notes"]
    assert [c.value for c in ws[3]] == [None] * 5
